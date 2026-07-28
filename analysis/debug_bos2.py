# -*- coding: utf-8 -*-
import openpyxl

SRC = r'K:\부서 공유\FI운용1부\10. 개인별폴더\김현수\rawdata\BOS3218_20260721.xlsx'
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]

rows = list(ws.iter_rows(values_only=True))
print('총 행:', len(rows))

# 01179가 종류형모펀드(F, idx5)로 나오는 구간과 그 앞뒤의 모펀드(B, idx1) 셀 표현
hits = [i for i, r in enumerate(rows) if str(r[5]).strip() in ('01179', '1179', '1179.0')]
print('01179 F열 등장 행:', hits[:5])
if hits:
    i0 = hits[0]
    for i in range(max(0, i0 - 4), min(len(rows), i0 + 4)):
        r = rows[i]
        print(i + 1, '| B=', repr(r[1]), '| F=', repr(r[5]), '| H=', repr(r[7]),
              '| 설정해지=', repr(r[3]))

# B열 값 유형 분포: None vs '' vs 값
import collections
kinds = collections.Counter()
for r in rows[1:]:
    v = r[1]
    if v is None:
        kinds['None'] += 1
    elif str(v).strip() == '':
        kinds['빈문자열'] += 1
    else:
        kinds['값'] += 1
print('B열 유형 분포:', dict(kinds))

# 빈문자열 B와 None B가 각각 어떤 맥락인지: 첫 빈문자열 행 주변
first_empty = next((i for i, r in enumerate(rows[1:], 1)
                    if r[1] is not None and str(r[1]).strip() == ''), None)
print('첫 빈문자열 B 행:', first_empty)
if first_empty:
    for i in range(max(0, first_empty - 3), min(len(rows), first_empty + 3)):
        r = rows[i]
        print(i + 1, '| B=', repr(r[1]), '| F=', repr(r[5]), '| H=', repr(r[7]))
